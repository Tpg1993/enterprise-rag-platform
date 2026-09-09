"""docx parser for Microsoft Word files."""

import hashlib
import io

import docx
from src.documents.base import DocumentParser, ParsedChunk, ParsedDocument


class DocxParser(DocumentParser):
    """Parses Microsoft Word (.docx) documents."""

    def parse(self, file_bytes: bytes, filename: str, document_id: str) -> ParsedDocument:
        checksum = hashlib.md5(file_bytes).hexdigest()
        doc = docx.Document(io.BytesIO(file_bytes))

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)

        chunks = [
            ParsedChunk(
                chunk_index=i,
                content=p,
                metadata={"filename": filename},
            )
            for i, p in enumerate(paragraphs)
        ]

        return ParsedDocument(
            document_id=document_id,
            filename=filename,
            file_type="docx",
            total_pages=1,
            full_text=full_text,
            chunks=chunks,
            checksum=checksum,
        )
